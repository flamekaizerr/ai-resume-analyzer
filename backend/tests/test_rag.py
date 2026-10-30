"""Tests for the rag module. They use a fake embedder so nothing gets downloaded."""

import hashlib
import re
import sys
from pathlib import Path

import numpy as np
import pytest

backend_dir = str(Path(__file__).resolve().parent.parent)
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from rag import qa
from rag.loaders import Chunk, chunk_documents, chunk_text, load_documents, read_bytes
from rag.store import FaissStore

SAMPLE_DIR = Path(backend_dir) / "rag" / "sample_docs"


class FakeEmbedder:
    """Bag of words hashed into a small vector. Same text always gives the same vector."""

    def __init__(self, dim=64):
        self.dim = dim

    def embed(self, texts):
        out = np.zeros((len(texts), self.dim), dtype="float32")
        for i, t in enumerate(texts):
            for word in re.findall(r"[a-z0-9]+", t.lower()):
                h = int(hashlib.md5(word.encode()).hexdigest(), 16)
                out[i, h % self.dim] += 1.0
        return out


@pytest.fixture
def docs(tmp_path):
    (tmp_path / "python.txt").write_text("Python FastAPI backend with PostgreSQL and Docker.")
    (tmp_path / "cooking.md").write_text("Recipe for tomato soup with basil and garlic.")
    (tmp_path / "music.txt").write_text("Guitar chords and piano scales for beginners.")
    (tmp_path / "ignore.csv").write_text("this,should,be,skipped")
    return tmp_path


def test_chunk_text_overlap():
    text = " ".join(f"w{i}" for i in range(300))
    chunks = chunk_text(text, chunk_size=100, overlap=30)
    assert len(chunks) > 1
    assert all(len(c) <= 100 for c in chunks)
    # end of one chunk should show up again at the start of the next
    for a, b in zip(chunks, chunks[1:]):
        assert a.split()[-1] in b.split()
    # nothing lost
    words = set()
    for c in chunks:
        words.update(c.split())
    assert words == set(text.split())


def test_chunk_text_short_and_empty():
    assert chunk_text("") == []
    assert chunk_text("   \n ") == []
    assert chunk_text("hello world", chunk_size=100, overlap=10) == ["hello world"]
    with pytest.raises(ValueError):
        chunk_text("abc", chunk_size=10, overlap=10)


def test_load_documents_skips_unsupported(docs):
    chunks = load_documents([str(docs)])
    assert sorted({c.source for c in chunks}) == ["cooking.md", "music.txt", "python.txt"]


def test_load_pdf(tmp_path):
    pypdf = pytest.importorskip("pypdf")
    # blank page pdf has no text, should give no chunks and not crash
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    path = tmp_path / "blank.pdf"
    with open(path, "wb") as f:
        writer.write(f)
    assert load_documents([str(path)]) == []


def test_store_save_load(tmp_path):
    emb = FakeEmbedder()
    chunks = [Chunk("python backend", "a.txt", 0), Chunk("tomato soup", "b.txt", 0)]
    store = FaissStore(emb.dim)
    store.add(chunks, emb.embed([c.text for c in chunks]))
    store.save(tmp_path / "idx")

    loaded = FaissStore.load(tmp_path / "idx")
    assert len(loaded) == 2
    assert loaded.chunks == chunks
    q = emb.embed(["python"])
    assert [r.chunk for r in loaded.search(q, k=2)] == [r.chunk for r in store.search(q, k=2)]


def test_load_missing_index(tmp_path):
    with pytest.raises(FileNotFoundError):
        FaissStore.load(tmp_path / "nope")


def test_retrieval_ranking(docs, tmp_path):
    emb = FakeEmbedder()
    store = qa.ingest([str(docs)], emb, folder=tmp_path / "idx")
    results = qa.retrieve(store, "python backend docker", emb, k=3)
    assert results[0].chunk.source == "python.txt"
    scores = [r.score for r in results]
    assert scores == sorted(scores, reverse=True)
    assert len(qa.retrieve(store, "soup", emb, k=10)) == 3


def test_fallback_answer_without_key(docs, tmp_path, monkeypatch):
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    emb = FakeEmbedder()
    store = qa.ingest([str(docs)], emb, folder=tmp_path / "idx")
    result = qa.answer("tomato soup recipe", store, emb, k=2)
    assert result["mode"] == "extractive"
    assert result["sources"][0]["source"] == "cooking.md"
    assert "[cooking.md #0]" in result["answer"]


def test_model_error_falls_back(docs, tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")

    def boom(*args, **kwargs):
        raise RuntimeError("network down")

    monkeypatch.setattr(qa, "model_answer", boom)
    emb = FakeEmbedder()
    store = qa.ingest([str(docs)], emb, folder=tmp_path / "idx")
    result = qa.answer("guitar", store, emb)
    assert result["mode"] == "extractive"
    assert "music.txt" in result["answer"]


def test_model_answer_used_when_key_set(docs, tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setattr(qa, "model_answer", lambda q, r: "You know Python [python.txt #0]")
    emb = FakeEmbedder()
    store = qa.ingest([str(docs)], emb, folder=tmp_path / "idx")
    result = qa.answer("python", store, emb)
    assert result["mode"] == "model"
    assert result["answer"] == "You know Python [python.txt #0]"


def test_cli_ingest_and_ask(docs, tmp_path, monkeypatch, capsys):
    import rag.embeddings
    from rag.__main__ import main

    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.setattr(rag.embeddings, "get_embedder", lambda: FakeEmbedder())
    idx = str(tmp_path / "idx")
    assert main(["--index-dir", idx, "ingest", str(docs)]) == 0
    assert main(["--index-dir", idx, "ask", "piano scales"]) == 0
    out = capsys.readouterr().out
    assert "Indexed 3 chunks" in out
    assert "[music.txt #0]" in out


def _model_available():
    try:
        from huggingface_hub import snapshot_download

        # only checks the local cache, no download
        snapshot_download("sentence-transformers/all-MiniLM-L6-v2", local_files_only=True)
        return True
    except Exception:
        return False


@pytest.mark.skipif(not _model_available(), reason="all-MiniLM-L6-v2 not downloaded")
def test_real_model_on_sample_docs(tmp_path, monkeypatch):
    from rag.embeddings import SentenceTransformerEmbedder

    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    emb = SentenceTransformerEmbedder("all-MiniLM-L6-v2")
    store = qa.ingest([str(SAMPLE_DIR)], emb, folder=tmp_path / "idx")
    results = qa.retrieve(store, "dashboards in Power BI or Tableau", emb, k=1)
    assert results[0].chunk.source == "jd_data_analyst.txt"


def test_model_answer_uses_openrouter(monkeypatch):
    import openai

    seen = {}

    class FakeClient:
        def __init__(self, **kwargs):
            seen.update(kwargs)
            self.chat = self
            self.completions = self

        def create(self, **kwargs):
            seen["model"] = kwargs["model"]
            msg = type("M", (), {"content": " ok [a.txt #0] "})
            choice = type("C", (), {"message": msg})
            return type("R", (), {"choices": [choice]})

    monkeypatch.setattr(openai, "OpenAI", FakeClient)
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.delenv("RAG_BASE_URL", raising=False)
    monkeypatch.delenv("RAG_MODEL", raising=False)

    from rag.store import SearchResult

    out = qa.model_answer("q", [SearchResult(Chunk("text", "a.txt", 0), 0.9)])
    assert out == "ok [a.txt #0]"
    assert seen["base_url"] == "https://openrouter.ai/api/v1"
    assert seen["api_key"] == "test-key"
    assert seen["model"] == "openai/gpt-4o-mini"


def test_load_docx(tmp_path):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_paragraph("Built a FastAPI service with PostgreSQL.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Docker Kubernetes"
    path = tmp_path / "resume.docx"
    d.save(str(path))

    chunks = load_documents([str(tmp_path)])
    assert [c.source for c in chunks] == ["resume.docx"]
    assert "FastAPI" in chunks[0].text
    assert "Kubernetes" in chunks[0].text


def test_read_bytes_unsupported():
    with pytest.raises(ValueError):
        read_bytes("notes.csv", b"a,b")


def test_chunk_documents_names():
    chunks = chunk_documents([("pasted_resume", "python dev"), ("pasted_jd", "needs python")])
    assert [(c.source, c.chunk_id) for c in chunks] == [("pasted_resume", 0), ("pasted_jd", 0)]


# route tests, using a small app with only the rag router

def _make_pdf(text):
    # tiny one page pdf with a text line, written by hand so we don't need another lib
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 200] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def _make_docx(text):
    import io

    import docx

    d = docx.Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def rag_client(tmp_path, monkeypatch):
    import httpx

    # same starlette/httpx workaround as test_agent.py
    orig = httpx.Client.__init__

    def patched(self, *args, **kwargs):
        kwargs.pop("app", None)
        return orig(self, *args, **kwargs)

    monkeypatch.setattr(httpx.Client, "__init__", patched)

    from fastapi import FastAPI
    from starlette.testclient import TestClient

    import rag.embeddings
    import rag.router as rag_router

    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.setenv("RAG_INDEX_DIR", str(tmp_path / "idx"))
    monkeypatch.setattr(rag.embeddings, "get_embedder", lambda: FakeEmbedder())
    monkeypatch.setattr(rag_router, "_sample_store", None)

    app = FastAPI()
    app.include_router(rag_router.router, prefix="/api")
    return TestClient(app)


def test_ask_route_with_pasted_text(rag_client):
    resp = rag_client.post("/api/rag/ask", data={
        "question": "kubernetes experience",
        "resume_text": "Worked on Kubernetes clusters and Helm charts.",
        "jd_text": "We need someone who knows accounting and payroll.",
        "k": "2",
    })
    assert resp.status_code == 200
    body = resp.json()
    assert body["index"] == "request"
    assert body["mode"] == "extractive"
    assert body["sources"][0]["source"] == "pasted_resume"
    assert "[pasted_resume #0]" in body["answer"]
    assert body["warning"] is None


def test_ask_route_warns_when_jd_is_missing(rag_client):
    resp = rag_client.post("/api/rag/ask", data={
        "question": "Which project matches the JD?",
        "resume_text": "Built Python projects with FastAPI.",
    })
    assert resp.status_code == 200
    assert "Add the JD" in resp.json()["warning"]


def test_ask_route_accepts_uploaded_jd_for_warning(rag_client):
    resp = rag_client.post(
        "/api/rag/ask",
        data={"question": "Which project matches the job description?"},
        files=[("files", ("backend_job.txt", b"Python role requirements", "text/plain"))],
    )
    assert resp.status_code == 200
    assert resp.json()["warning"] is None


@pytest.mark.parametrize("name,maker", [
    ("resume.txt", lambda t: t.encode()),
    ("resume.md", lambda t: ("# Resume\n" + t).encode()),
    ("resume.pdf", _make_pdf),
    ("resume.docx", _make_docx),
])
def test_ask_route_with_upload(rag_client, name, maker):
    if name.endswith(".docx"):
        pytest.importorskip("docx")
    data = maker("Built Terraform modules for AWS networking")
    resp = rag_client.post(
        "/api/rag/ask",
        data={"question": "terraform aws", "jd_text": "Cooking and baking skills wanted."},
        files=[("files", (name, data, "application/octet-stream"))],
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["sources"][0]["source"] == name
    assert "Terraform" in body["sources"][0]["text"]


def test_ask_route_rejects_bad_files(rag_client):
    resp = rag_client.post("/api/rag/ask", data={"question": "x"},
                           files=[("files", ("data.csv", b"a,b", "text/csv"))])
    assert resp.status_code == 400
    resp = rag_client.post("/api/rag/ask", data={"question": "x"},
                           files=[("files", ("broken.pdf", b"not a pdf", "application/pdf"))])
    assert resp.status_code == 400


def test_ask_route_empty_question(rag_client):
    assert rag_client.post("/api/rag/ask", data={"question": "  "}).status_code == 400


def test_ask_route_falls_back_to_samples_then_saved(rag_client, docs, tmp_path):
    body = rag_client.post("/api/rag/ask", data={"question": "Power BI dashboards"}).json()
    assert body["index"] == "sample"
    assert body["sources"]

    qa.ingest([str(docs)], FakeEmbedder(), folder=tmp_path / "idx")
    body = rag_client.post("/api/rag/ask", data={"question": "guitar"}).json()
    assert body["index"] == "saved"
    assert body["sources"][0]["source"] == "music.txt"


def test_jd_warning_rule():
    q = "Which project from my resume matches the JD?"
    assert qa.jd_warning(q, ["resume.pdf"])
    assert qa.jd_warning(q, ["pasted_resume"])
    assert qa.jd_warning(q, ["resume.pdf", "backend_jd.txt"]) is None
    assert qa.jd_warning(q, ["pasted_resume", "pasted_jd"]) is None
    assert qa.jd_warning("What JDs fit me?", ["resume.pdf"])
    assert qa.jd_warning("Which of my projects use Docker?", ["resume.pdf"]) is None


def test_no_jd_warning_on_sample_fallback(rag_client):
    body = rag_client.post("/api/rag/ask", data={"question": "Which project matches the JD?"}).json()
    assert body["index"] == "sample"
    assert body["warning"] is None
